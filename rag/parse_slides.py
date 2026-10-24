"""Parse teacher slide decks (PDF export or .pptx) into chunks + slide images.

Consecutive slides that share a title (e.g. a worked example spread across 8
slides) are grouped, then packed into chunks of up to ~CHUNK_MAX_CHARS. Every
chunk remembers its slide range, and each slide is rendered to an image so the
tutor can show the teacher's own diagrams / Gantt charts next to an answer.
"""
from __future__ import annotations

import re
from pathlib import Path

import pymupdf

import json

from .config import ANNOTATIONS_DIR, CHUNK_MAX_CHARS
from .textutil import clean_lines, clean_text

VISUAL_TEXT_LIMIT = 150  # slides with less text than this are mostly pictures


def _slide_title(lines: list[str]) -> str:
    for l in lines:
        l = l.strip()
        if l:
            return re.sub(r"\s+", " ", l)
    return ""


def _group_key(title: str) -> str:
    return re.sub(r"[^a-z]", "", title.lower())


def _read_pdf_slides(path: Path, fig_dir: Path, src: dict):
    doc = pymupdf.open(path)
    skip = set(src.get("skip_pages", []))
    slides = []
    for i, page in enumerate(doc, start=1):
        if i in skip:
            continue
        raw = page.get_text()
        lines = [l for l in raw.splitlines() if l.strip()]
        title = _slide_title(lines)
        body = clean_text("\n".join(lines[1:]))
        # Render slide. JPEG keeps the 200-slide deck small.
        fid = f"{src['id']}_slide_{i:03d}"
        out = fig_dir / f"{fid}.jpg"
        zoom = 1000 / page.rect.width
        page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom)).save(out, jpg_quality=72)
        slides.append(
            {
                "no": i,
                "title": title,
                "text": body,
                "image": f"figures/{out.name}",
                "fig_id": fid,
                "n_images": len(page.get_images()),
            }
        )
    return slides


def _read_pptx_slides(path: Path, fig_dir: Path, src: dict):
    """Text (line breaks kept, for code), tables, notes, and the slide's first picture.

    Rendering whole .pptx slides needs LibreOffice, so a slide's image is its first picture."""
    from pptx import Presentation  # optional dependency

    prs = Presentation(str(path))
    slides = []
    for i, s in enumerate(prs.slides, start=1):
        if i in set(src.get("skip_pages", [])):
            continue
        texts, pictures = [], []
        for shp in s.shapes:
            if shp.has_text_frame:
                texts.append(shp.text_frame.text)
            if getattr(shp, "has_table", False) and shp.has_table:
                for row in shp.table.rows:
                    texts.append(" | ".join(c.text.replace("\n", " ") for c in row.cells))
            if shp.shape_type == 13:  # picture
                pictures.append(shp.image)
        lines = [l for t in texts for l in t.splitlines() if l.strip()]
        notes = ""
        if s.has_notes_slide:
            notes = s.notes_slide.notes_text_frame.text
        image = fid = None
        if pictures:
            fid = f"{src['id']}_slide_{i:03d}"
            out = fig_dir / f"{fid}.{pictures[0].ext}"
            out.write_bytes(pictures[0].blob)
            image = f"figures/{out.name}"
        slides.append(
            {
                "no": i,
                "title": _slide_title(lines),
                "text": clean_lines("\n".join(lines[1:] + ([notes] if notes else []))),
                "image": image,
                "fig_id": fid,
                "n_images": len(pictures),
            }
        )
    return slides


def parse_slides(src: dict, root: Path, fig_dir: Path):
    fig_dir.mkdir(parents=True, exist_ok=True)
    path = root / src["file"]
    if path.suffix.lower() == ".pptx":
        slides = _read_pptx_slides(path, fig_dir, src)
    else:
        slides = _read_pdf_slides(path, fig_dir, src)

    # ---- merge hand transcriptions of image-only slides ---------------------
    ann_path = ANNOTATIONS_DIR / f"{src['id']}.json"
    if ann_path.exists():
        ann = json.loads(ann_path.read_text(encoding="utf-8"))
        for s in slides:
            a = ann.get(str(s["no"]))
            if not a:
                continue
            extra = f"[Transcribed from slide image] {a['text']}"
            if a.get("note"):
                extra += f" [Knowledge-base note - confirm with faculty] {a['note']}"
            s["text"] = (s["text"] + " " + extra).strip()
            s["n_images"] = max(s["n_images"], 1)

    # ---- group consecutive slides with the same title -------------------------
    groups: list[list[dict]] = []
    for s in slides:
        if groups and _group_key(groups[-1][0]["title"]) == _group_key(s["title"]):
            groups[-1].append(s)
        else:
            groups.append([s])

    chunks, figures = [], []
    for g in groups:
        title = g[0]["title"]
        cur: list[dict] = []

        packs: list[list[dict]] = []
        size = 0
        for s in g:
            if cur and (size + len(s["text"]) > CHUNK_MAX_CHARS or len(cur) >= 8):
                packs.append(cur)
                cur, size = [], 0
            cur.append(s)
            size += len(s["text"]) + 20
        if cur:
            packs.append(cur)

        for p in packs:
            a, b = p[0]["no"], p[-1]["no"]
            body = "\n".join(
                f"[Slide {s['no']}] {s['text']}" if s["text"] else f"[Slide {s['no']}] (diagram)"
                for s in p
            )
            chunks.append(
                {
                    "source_id": src["id"],
                    "source_kind": src["kind"],
                    "source_short": src["short"],
                    "chapter": "",
                    "section": title,
                    "kind": "example" if re.search(r"example|working", body + title, re.I) else "explanation",
                    "page_start": str(a),
                    "page_end": str(b),
                    "pdf_page_start": a,
                    "fig_refs": [s["fig_id"] for s in p if s["fig_id"]],
                    "text": f"{title}\n{body}",
                }
            )

    for s in slides:
        if not s["fig_id"]:
            continue
        figures.append(
            {
                "id": s["fig_id"],
                "source_id": src["id"],
                "source_kind": src["kind"],
                "source_short": src["short"],
                "label": f"Slide {s['no']}",
                "caption": s["title"],
                "chapter": "",
                "section": s["title"],
                "page": str(s["no"]),
                "pdf_page": s["no"],
                "image": s["image"],
                "context": s["text"][:700],
                "visual": s["n_images"] > 0 or len(s["text"]) < VISUAL_TEXT_LIMIT,
            }
        )
    return chunks, figures
